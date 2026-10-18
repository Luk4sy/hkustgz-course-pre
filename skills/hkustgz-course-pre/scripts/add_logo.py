"""Put a logo in the top-right corner of every slide.

Usage: python add_logo.py in.pptx out.pptx --color logo_color.png --white logo_white.png [--width 1.95] [--dark 1,6]
With LibreOffice (soffice) and pdftoppm installed, each slide's background brightness is measured:
light corner -> color logo, dark corner -> white logo. Without them, or with --dark, the slides
listed in --dark (1-based) get the white logo and all others the color logo. Titles that would overlap are nudged left,
pictures that would overlap are shrunk.
"""
import argparse, glob, os, shutil, subprocess, tempfile
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ap = argparse.ArgumentParser()
ap.add_argument("src"); ap.add_argument("dst")
ap.add_argument("--color", required=True); ap.add_argument("--white", required=True)
ap.add_argument("--width", type=float, default=1.95, help="logo width in inches")
ap.add_argument("--dark", default=None, help="comma-separated slide numbers with dark backgrounds (skips measuring)")
a = ap.parse_args()

soffice = shutil.which("soffice") or shutil.which("libreoffice") or \
    ("/Applications/LibreOffice.app/Contents/MacOS/soffice" if os.path.exists("/Applications/LibreOffice.app/Contents/MacOS/soffice") else None)
pages = None
if a.dark is None and soffice and shutil.which("pdftoppm"):
    tmp = tempfile.mkdtemp()
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, a.src], check=True, capture_output=True)
    pdf = glob.glob(os.path.join(tmp, "*.pdf"))[0]
    subprocess.run(["pdftoppm", "-png", "-r", "40", pdf, os.path.join(tmp, "p")], check=True)
    pages = sorted(glob.glob(os.path.join(tmp, "p-*.png")))
elif a.dark is None:
    print("LibreOffice/pdftoppm not found: using the color logo on every slide (pass --dark 1,6 for dark slides)")
dark = {int(x) for x in (a.dark or "").split(",") if x.strip()}

lw_px, lh_px = Image.open(a.color).size
p = Presentation(a.src)
PW, PH = p.slide_width, p.slide_height
lw = Inches(a.width); lh = int(lw * lh_px / lw_px)
px = PW - Inches(0.45) - lw; py = Inches(0.2)
limit = px - Inches(0.2); bottom = py + lh + Inches(0.12)

modes = []
for idx, s in enumerate(p.slides, 1):
    if pages:
        im = np.asarray(Image.open(pages[idx - 1]).convert("L")).astype(float); H, W = im.shape
        region = im[int(H * py / PH):int(H * (py + lh) / PH), int(W * px / PW):int(W * (px + lw) / PW)]
        light = region.mean() > 170
    else:
        light = idx not in dark
    modes.append("C" if light else "W")
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.top < Inches(0.7) \
                and sh.left + sh.width > limit and sh.left < px:
            sh.left = max(0, limit - sh.width)
        elif sh.shape_type == 13 and sh.name != "logo" and sh.top < bottom and sh.left + sh.width > px \
                and sh.top + sh.height > py and sh.width < PW * 0.9:
            b = sh.top + sh.height; r = sh.left + sh.width
            nh = b - bottom; nw = int(sh.width * nh / sh.height)
            if nh < sh.height * 0.3:  # picture sits almost entirely under the logo: leave it
                continue
            sh.top, sh.height, sh.width, sh.left = bottom, nh, nw, r - nw
    pic = s.shapes.add_picture(a.color if light else a.white, px, py, lw, lh)
    pic.name = "logo"
p.save(a.dst)
print("per-slide logo (C=color, W=white):", "".join(modes))
